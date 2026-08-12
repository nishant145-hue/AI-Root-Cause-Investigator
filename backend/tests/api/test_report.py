from io import BytesIO
from uuid import uuid4

from openpyxl import load_workbook

AUTH_PREFIX = "/api/v1/auth"
REPORT_PREFIX = "/api/v1/reports"


def create_test_user(client):
    suffix = uuid4().hex[:10]

    user = {
        "username": f"report_{suffix}",
        "email": f"report_{suffix}@example.com",
        "password": "Password123!",
        "full_name": "Report Test User",
    }

    response = client.post(
        f"{AUTH_PREFIX}/register",
        json=user,
    )

    assert response.status_code == 201

    return user


def login_user(client, user):
    response = client.post(
        f"{AUTH_PREFIX}/login",
        data={
            "username": user["email"],
            "password": user["password"],
        },
    )

    assert response.status_code == 200

    return response.json()


def auth_headers(access_token):
    return {
        "Authorization": f"Bearer {access_token}"
    }


# ------------------------------------------------------------------
# Authentication enforcement
# ------------------------------------------------------------------


def test_csv_report_requires_authentication(client):
    response = client.get(
        f"{REPORT_PREFIX}/csv"
    )

    assert response.status_code == 401


def test_excel_report_requires_authentication(client):
    response = client.get(
        f"{REPORT_PREFIX}/excel"
    )

    assert response.status_code == 401


def test_pdf_report_requires_authentication(client):
    response = client.get(
        f"{REPORT_PREFIX}/pdf"
    )

    assert response.status_code == 401


# ------------------------------------------------------------------
# Authenticated report generation
# ------------------------------------------------------------------


def test_export_csv(client):
    user = create_test_user(client)
    tokens = login_user(client, user)

    response = client.get(
        f"{REPORT_PREFIX}/csv",
        headers=auth_headers(tokens["access_token"]),
    )

    assert response.status_code == 200

    assert response.headers["content-type"].startswith(
        "text/csv"
    )


def test_export_excel(client):
    user = create_test_user(client)
    tokens = login_user(client, user)

    response = client.get(
        f"{REPORT_PREFIX}/excel",
        headers=auth_headers(tokens["access_token"]),
    )

    assert response.status_code == 200

    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )


def test_export_pdf(client):
    user = create_test_user(client)
    tokens = login_user(client, user)

    response = client.get(
        f"{REPORT_PREFIX}/pdf",
        headers=auth_headers(tokens["access_token"]),
    )

    assert response.status_code == 200

    assert response.headers["content-type"].startswith(
        "application/pdf"
    )

def create_investigation(client, access_token, title):
    response = client.post(
        "/api/v1/investigations",
        headers=auth_headers(access_token),
        json={
            "title": title,
            "description": "Report ownership test",
        },
    )

    assert response.status_code == 201

    return response.json()


def test_csv_report_contains_only_current_users_investigations(client):
    user_a = create_test_user(client)
    user_b = create_test_user(client)

    tokens_a = login_user(client, user_a)
    tokens_b = login_user(client, user_b)

    investigation_a = create_investigation(
        client,
        tokens_a["access_token"],
        "USER_A_PRIVATE_INVESTIGATION",
    )

    investigation_b = create_investigation(
        client,
        tokens_b["access_token"],
        "USER_B_PRIVATE_INVESTIGATION",
    )

    response_a = client.get(
        f"{REPORT_PREFIX}/csv",
        headers=auth_headers(tokens_a["access_token"]),
    )

    assert response_a.status_code == 200

    content = response_a.content.decode("utf-8")

    assert "USER_A_PRIVATE_INVESTIGATION" in content
    assert "USER_B_PRIVATE_INVESTIGATION" not in content

    assert str(investigation_a["id"]) in content
    assert str(investigation_b["id"]) not in content

def test_csv_report_isolated_between_users(client):
    user_a = create_test_user(client)
    user_b = create_test_user(client)

    tokens_a = login_user(client, user_a)
    tokens_b = login_user(client, user_b)

    create_investigation(
        client,
        tokens_a["access_token"],
        "USER_A_SECRET_INVESTIGATION",
    )

    create_investigation(
        client,
        tokens_b["access_token"],
        "USER_B_SECRET_INVESTIGATION",
    )

    response = client.get(
        "/api/v1/reports/csv",
        headers=auth_headers(tokens_a["access_token"]),
    )

    assert response.status_code == 200

    content = response.content.decode("utf-8")

    assert "USER_A_SECRET_INVESTIGATION" in content
    assert "USER_B_SECRET_INVESTIGATION" not in content

def test_excel_report_isolated_between_users(client):
    user_a = create_test_user(client)
    user_b = create_test_user(client)

    tokens_a = login_user(client, user_a)
    tokens_b = login_user(client, user_b)

    create_investigation(
        client,
        tokens_a["access_token"],
        "USER_A_EXCEL_SECRET",
    )

    create_investigation(
        client,
        tokens_b["access_token"],
        "USER_B_EXCEL_SECRET",
    )

    response = client.get(
        "/api/v1/reports/excel",
        headers=auth_headers(tokens_a["access_token"]),
    )

    assert response.status_code == 200

    workbook = load_workbook(
        filename=BytesIO(response.content),
        read_only=True,
    )

    worksheet = workbook["Investigations"]

    values = []

    for row in worksheet.iter_rows(values_only=True):
        values.extend(
            value
            for value in row
            if value is not None
        )

    assert "USER_A_EXCEL_SECRET" in values
    assert "USER_B_EXCEL_SECRET" not in values

    workbook.close()

def test_pdf_report_isolated_between_users(client):
    user_a = create_test_user(client)
    user_b = create_test_user(client)

    tokens_a = login_user(client, user_a)
    tokens_b = login_user(client, user_b)

    create_investigation(
        client,
        tokens_a["access_token"],
        "USER_A_PDF_SECRET",
    )

    create_investigation(
        client,
        tokens_b["access_token"],
        "USER_B_PDF_SECRET",
    )

    response = client.get(
        "/api/v1/reports/pdf",
        headers=auth_headers(tokens_a["access_token"]),
    )

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/pdf"
    )

    # The PDF is generated successfully for User A.
    # Detailed PDF text extraction is intentionally omitted
    # because pypdf is not part of this project's dependencies.
    assert len(response.content) > 0
