from datetime import datetime
from pathlib import Path
from uuid import uuid4

import pandas as pd
from app.repositories.dashboard_repository import DashboardRepository
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from sqlmodel import Session


class ReportService:
    """
    Service responsible for generating dashboard reports.

    Security responsibilities:
    - Report data is always scoped to an authenticated user.
    - Spreadsheet exports protect against formula injection.
    - Generated report filenames are unique to avoid collisions.
    """

    def __init__(self, session: Session):
        self.repository = DashboardRepository(session)

    def get_report_data(self, user_id: int):
        """
        Return investigation data belonging only to the supplied user.
        """

        if user_id <= 0:
            raise ValueError("Invalid user ID.")

        investigations = self.repository.get_report_data(
            user_id=user_id,
        )

        report_data = []

        for investigation in investigations:
            report_data.append(
                {
                    "id": investigation.id,
                    "title": investigation.title,
                    "status": (
                        investigation.status.value
                        if investigation.status
                        else None
                    ),
                    "severity": investigation.severity,
                    "confidence": investigation.confidence,
                    "root_cause": investigation.root_cause,
                    "failed_component": investigation.failed_component,
                    "created_at": (
                        investigation.created_at.strftime(
                            "%Y-%m-%d %H:%M:%S"
                        )
                        if investigation.created_at
                        else None
                    ),
                }
            )

        return report_data

    @staticmethod
    def _sanitize_spreadsheet_value(value):
        """
        Prevent spreadsheet formula injection.

        Spreadsheet applications may interpret strings beginning with
        =, +, -, or @ as formulas.

        Prefixing such values with an apostrophe forces spreadsheet
        applications to treat them as text.

        Numeric values are left untouched.
        """

        if not isinstance(value, str):
            return value

        stripped = value.lstrip()

        if stripped.startswith(("=", "+", "-", "@")):
            return "'" + value

        return value

    def _create_report_path(self, extension: str) -> Path:
        """
        Create a unique report path.

        Unique filenames prevent simultaneous users from overwriting
        each other's generated reports.
        """

        reports_dir = Path("reports")
        reports_dir.mkdir(
            exist_ok=True,
        )

        filename = (
            f"investigations_"
            f"{uuid4().hex}"
            f"{extension}"
        )

        return reports_dir / filename

    def export_csv(self, user_id: int) -> Path:
        """
        Export the authenticated user's investigations as CSV.
        """

        report_data = self.get_report_data(user_id)

        file_path = self._create_report_path(".csv")

        sanitized_data = []

        for investigation in report_data:
            sanitized_data.append(
                {
                    key: self._sanitize_spreadsheet_value(value)
                    for key, value in investigation.items()
                }
            )

        dataframe = pd.DataFrame(sanitized_data)

        dataframe.to_csv(
            file_path,
            index=False,
        )

        return file_path

    def export_excel(self, user_id: int) -> Path:
        """
        Export the authenticated user's investigations as Excel.
        """

        report_data = self.get_report_data(user_id)

        file_path = self._create_report_path(".xlsx")

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Investigations"

        headers = [
            "ID",
            "Title",
            "Status",
            "Severity",
            "Confidence",
            "Root Cause",
            "Failed Component",
            "Created At",
        ]

        # Header row
        for column, header in enumerate(
            headers,
            start=1,
        ):
            cell = worksheet.cell(
                row=1,
                column=column,
            )
            cell.value = header
            cell.font = Font(
                bold=True,
            )

        # Investigation data
        for row_index, investigation in enumerate(
            report_data,
            start=2,
        ):
            worksheet.cell(
                row=row_index,
                column=1,
            ).value = investigation["id"]

            worksheet.cell(
                row=row_index,
                column=2,
            ).value = self._sanitize_spreadsheet_value(
                investigation["title"]
            )

            worksheet.cell(
                row=row_index,
                column=3,
            ).value = self._sanitize_spreadsheet_value(
                investigation["status"]
            )

            worksheet.cell(
                row=row_index,
                column=4,
            ).value = self._sanitize_spreadsheet_value(
                investigation["severity"]
            )

            worksheet.cell(
                row=row_index,
                column=5,
            ).value = investigation["confidence"]

            worksheet.cell(
                row=row_index,
                column=6,
            ).value = self._sanitize_spreadsheet_value(
                investigation["root_cause"]
            )

            worksheet.cell(
                row=row_index,
                column=7,
            ).value = self._sanitize_spreadsheet_value(
                investigation["failed_component"]
            )

            worksheet.cell(
                row=row_index,
                column=8,
            ).value = self._sanitize_spreadsheet_value(
                investigation["created_at"]
                if investigation["created_at"]
                else ""
            )

        # Auto-size columns
        for column_cells in worksheet.columns:
            length = max(
                len(str(cell.value))
                if cell.value is not None
                else 0
                for cell in column_cells
            )

            worksheet.column_dimensions[
                get_column_letter(
                    column_cells[0].column
                )
            ].width = min(
                length + 3,
                50,
            )

        workbook.save(file_path)

        return file_path

    def export_pdf(self, user_id: int) -> Path:
        """
        Export the authenticated user's investigations as PDF.
        """

        report_data = self.get_report_data(user_id)

        file_path = self._create_report_path(".pdf")

        document = SimpleDocTemplate(
            str(file_path),
        )

        styles = getSampleStyleSheet()

        elements = []

        elements.append(
            Paragraph(
                "<b>AI Root Cause Investigator</b>",
                styles["Title"],
            )
        )

        elements.append(
            Paragraph(
                "Investigation Report",
                styles["Heading2"],
            )
        )

        elements.append(
            Spacer(
                1,
                0.25 * inch,
            )
        )

        elements.append(
            Paragraph(
                f"Generated: "
                f"{datetime.now():%Y-%m-%d %H:%M:%S}",
                styles["Normal"],
            )
        )

        elements.append(
            Spacer(
                1,
                0.25 * inch,
            )
        )

        elements.append(
            Paragraph(
                f"Total Investigations: "
                f"{len(report_data)}",
                styles["Heading3"],
            )
        )

        elements.append(
            Spacer(
                1,
                0.25 * inch,
            )
        )

        table_data = [
            [
                "ID",
                "Title",
                "Status",
                "Severity",
                "Confidence",
            ]
        ]

        for item in report_data:
            table_data.append(
                [
                    str(item["id"]),
                    str(item["title"]),
                    str(item["status"]),
                    str(item["severity"]),
                    str(item["confidence"]),
                ]
            )

        table = Table(
            table_data,
            repeatRows=1,
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.darkblue,
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        1,
                        colors.black,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "ALIGN",
                        (0, 0),
                        (-1, -1),
                        "CENTER",
                    ),
                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, 0),
                        8,
                    ),
                    (
                        "BACKGROUND",
                        (0, 1),
                        (-1, -1),
                        colors.beige,
                    ),
                ]
            )
        )

        elements.append(table)

        document.build(elements)

        return file_path