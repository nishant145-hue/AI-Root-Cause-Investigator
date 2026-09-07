from pathlib import Path
from uuid import uuid4

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


class AnalyticsReportService:
    """
    Generates downloadable reports from investigation analytics.

    Supported formats:
    - JSON-compatible dictionary
    - CSV
    - Excel
    - PDF
    """

    def __init__(self, analytics: dict):
        self.analytics = analytics or {}

    # ---------------------------------------------------------
    # Helpers
    # ---------------------------------------------------------

    def _create_report_path(
        self,
        extension: str,
    ) -> Path:
        reports_dir = Path("reports")
        reports_dir.mkdir(
            exist_ok=True,
        )

        return (
            reports_dir
            / (
                f"investigation_analytics_"
                f"{uuid4().hex}"
                f"{extension}"
            )
        )

    @staticmethod
    def _sanitize_spreadsheet_value(
        value,
    ):
        """
        Prevent spreadsheet formula injection.
        """

        if not isinstance(value, str):
            return value

        stripped = value.lstrip()

        if stripped.startswith(
            ("=", "+", "-", "@")
        ):
            return "'" + value

        return value

    # ---------------------------------------------------------
    # JSON
    # ---------------------------------------------------------

    def export_json_data(self) -> dict:
        """
        Return analytics in a JSON-serializable structure.
        """

        return self.analytics.copy()

    # ---------------------------------------------------------
    # Flatten analytics
    # ---------------------------------------------------------

    def _overview_rows(self) -> list[dict]:
        overview = self.analytics.get(
            "overview",
            {},
        )

        return [
            {
                "metric": "Total Executions",
                "value": overview.get(
                    "total_executions",
                    0,
                ),
            },
            {
                "metric": "Successful Executions",
                "value": overview.get(
                    "successful_executions",
                    0,
                ),
            },
            {
                "metric": "Failed Executions",
                "value": overview.get(
                    "failed_executions",
                    0,
                ),
            },
            {
                "metric": "Total Duration (ms)",
                "value": overview.get(
                    "total_duration_ms",
                    0.0,
                ),
            },
            {
                "metric": "Retry Count",
                "value": overview.get(
                    "retry_count",
                    0,
                ),
            },
            {
                "metric": "Efficiency Score",
                "value": overview.get(
                    "efficiency_score",
                    0.0,
                ),
            },
        ]

    def _agent_rows(self) -> list[dict]:
        return self.analytics.get(
            "agent_performance",
            [],
        )

    def _bottleneck_rows(self) -> list[dict]:
        return self.analytics.get(
            "bottlenecks",
            [],
        )

    def _failure_rows(self) -> list[dict]:
        metrics = self.analytics.get(
            "failure_retry_metrics",
            {},
        )

        rows = []

        failed_agents = metrics.get(
            "failed_agents",
            {},
        )

        for agent, count in failed_agents.items():
            rows.append(
                {
                    "agent": agent,
                    "failures": count,
                    "retries": metrics.get(
                        "retry_by_agent",
                        {},
                    ).get(
                        agent,
                        0,
                    ),
                }
            )

        return rows

    def _timeline_rows(self) -> list[dict]:
        return self.analytics.get(
            "timeline",
            [],
        )

    # ---------------------------------------------------------
    # CSV
    # ---------------------------------------------------------

    def export_csv(self) -> Path:
        """
        Export analytics as CSV.

        Multiple sections are written into one CSV file.
        """

        file_path = self._create_report_path(
            ".csv"
        )

        sections = []

        sections.append(
            pd.DataFrame(
                self._overview_rows()
            )
        )

        agent_rows = self._agent_rows()

        if agent_rows:
            sections.append(
                pd.DataFrame(agent_rows)
            )

        bottleneck_rows = (
            self._bottleneck_rows()
        )

        if bottleneck_rows:
            sections.append(
                pd.DataFrame(
                    bottleneck_rows
                )
            )

        failure_rows = self._failure_rows()

        if failure_rows:
            sections.append(
                pd.DataFrame(
                    failure_rows
                )
            )

        timeline_rows = (
            self._timeline_rows()
        )

        if timeline_rows:
            sections.append(
                pd.DataFrame(
                    timeline_rows
                )
            )

        with open(
            file_path,
            "w",
            encoding="utf-8",
            newline="",
        ) as file:
            for index, dataframe in enumerate(
                sections
            ):
                if index > 0:
                    file.write("\n")

                dataframe = dataframe.map(
                    self._sanitize_spreadsheet_value
                )

                dataframe.to_csv(
                    file,
                    index=False,
                )

        return file_path

    # ---------------------------------------------------------
    # Excel
    # ---------------------------------------------------------

    def export_excel(self) -> Path:
        """
        Export analytics into a structured Excel workbook.
        """

        file_path = self._create_report_path(
            ".xlsx"
        )

        workbook = Workbook()

        overview_sheet = workbook.active
        overview_sheet.title = "Overview"

        self._write_sheet(
            overview_sheet,
            self._overview_rows(),
        )

        agent_sheet = workbook.create_sheet(
            "Agent Performance"
        )

        self._write_sheet(
            agent_sheet,
            self._agent_rows(),
        )

        bottleneck_sheet = (
            workbook.create_sheet(
                "Bottlenecks"
            )
        )

        self._write_sheet(
            bottleneck_sheet,
            self._bottleneck_rows(),
        )

        failure_sheet = (
            workbook.create_sheet(
                "Failures & Retries"
            )
        )

        self._write_sheet(
            failure_sheet,
            self._failure_rows(),
        )

        timeline_sheet = (
            workbook.create_sheet(
                "Timeline"
            )
        )

        self._write_sheet(
            timeline_sheet,
            self._timeline_rows(),
        )

        workbook.save(file_path)

        return file_path

    def _write_sheet(
        self,
        worksheet,
        rows: list[dict],
    ):
        if not rows:
            worksheet.cell(
                row=1,
                column=1,
            ).value = "No data"

            return

        headers = list(rows[0].keys())

        for column, header in enumerate(
            headers,
            start=1,
        ):
            cell = worksheet.cell(
                row=1,
                column=column,
            )

            cell.value = header
            cell.font = Font(
                bold=True
            )

        for row_index, row in enumerate(
            rows,
            start=2,
        ):
            for column, header in enumerate(
                headers,
                start=1,
            ):
                value = row.get(
                    header
                )

                worksheet.cell(
                    row=row_index,
                    column=column,
                ).value = (
                    self._sanitize_spreadsheet_value(
                        value
                    )
                )

        for column_cells in worksheet.columns:
            length = max(
                len(
                    str(cell.value)
                )
                if cell.value is not None
                else 0
                for cell in column_cells
            )

            worksheet.column_dimensions[
                get_column_letter(
                    column_cells[0].column
                )
            ].width = min(
                length + 3,
                50,
            )

    # ---------------------------------------------------------
    # PDF
    # ---------------------------------------------------------

    def export_pdf(self) -> Path:
        """
        Generate a human-readable analytics PDF report.
        """

        file_path = self._create_report_path(
            ".pdf"
        )

        document = SimpleDocTemplate(
            str(file_path),
            pagesize=landscape(letter),
            rightMargin=0.4 * inch,
            leftMargin=0.4 * inch,
            topMargin=0.4 * inch,
            bottomMargin=0.4 * inch,
        )

        styles = getSampleStyleSheet()

        elements = []

        investigation_id = self.analytics.get(
            "investigation_id",
            "Unknown",
        )

        elements.append(
            Paragraph(
                "<b>AI Root Cause Investigator</b>",
                styles["Title"],
            )
        )

        elements.append(
            Paragraph(
                "Investigation Analytics Report",
                styles["Heading2"],
            )
        )

        elements.append(
            Spacer(
                1,
                0.2 * inch,
            )
        )

        elements.append(
            Paragraph(
                f"Investigation ID: "
                f"{investigation_id}",
                styles["Normal"],
            )
        )

        elements.append(
            Spacer(
                1,
                0.2 * inch,
            )
        )

        # Overview
        elements.append(
            Paragraph(
                "Execution Overview",
                styles["Heading3"],
            )
        )

        overview_rows = self._overview_rows()

        overview_table_data = [
            ["Metric", "Value"]
        ]

        for row in overview_rows:
            overview_table_data.append(
                [
                    str(row["metric"]),
                    str(row["value"]),
                ]
            )

        elements.append(
            self._create_table(
                overview_table_data
            )
        )

        elements.append(
            Spacer(
                1,
                0.2 * inch,
            )
        )

        # Agent performance
        elements.append(
            Paragraph(
                "Agent Performance",
                styles["Heading3"],
            )
        )

        agent_rows = self._agent_rows()

        if agent_rows:
            headers = list(
                agent_rows[0].keys()
            )

            data = [
                headers
            ]

            for row in agent_rows:
                data.append(
                    [
                        str(
                            row.get(
                                header,
                                "",
                            )
                        )
                        for header in headers
                    ]
                )

            elements.append(
                self._create_table(data)
            )
        else:
            elements.append(
                Paragraph(
                    "No agent performance data.",
                    styles["Normal"],
                )
            )

        elements.append(
            Spacer(
                1,
                0.2 * inch,
            )
        )

        # Bottlenecks
        elements.append(
            Paragraph(
                "Bottlenecks",
                styles["Heading3"],
            )
        )

        bottleneck_rows = (
            self._bottleneck_rows()
        )

        if bottleneck_rows:
            headers = list(
                bottleneck_rows[0].keys()
            )

            data = [
                headers
            ]

            for row in bottleneck_rows:
                data.append(
                    [
                        str(
                            row.get(
                                header,
                                "",
                            )
                        )
                        for header in headers
                    ]
                )

            elements.append(
                self._create_table(data)
            )
        else:
            elements.append(
                Paragraph(
                    "No bottlenecks detected.",
                    styles["Normal"],
                )
            )

        document.build(elements)

        return file_path

    @staticmethod
    def _create_table(
        data: list[list],
    ) -> Table:
        table = Table(
            data,
            repeatRows=1,
        )

        table.setStyle(
            TableStyle(
                [
                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.darkblue,
                    ),
                    (
                        "TEXTCOLOR",
                        (0, 0),
                        (-1, 0),
                        colors.white,
                    ),
                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold",
                    ),
                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.black,
                    ),
                    (
                        "ALIGN",
                        (0, 0),
                        (-1, -1),
                        "CENTER",
                    ),
                    (
                        "VALIGN",
                        (0, 0),
                        (-1, -1),
                        "MIDDLE",
                    ),
                ]
            )
        )

        return table
